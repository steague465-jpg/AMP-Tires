#!/usr/bin/env python3
"""
AMP tire competitor price scraper (shop.app + each store's own Shopify site)

What it does
  1. Searches shop.app for every AMP model (Terrain Pro A/T P, Terrain Attack A/T A,
     Mud Terrain Attack M/T, R/T) across rim sizes, scrolling through every result.
  2. Figures out each store's real website domain.
  3. Pulls every AMP tire listing straight from each store's Shopify catalog
     (all sizes they carry, not just what shop.app happened to show).
  4. Checks real shipping cost to your zip at qty 1 and qty 4 using the store's cart.
  5. Writes amp_prices.xlsx with a "Beat Price" sheet (lowest landed price per model + size).

Setup (one time)
  pip install playwright requests openpyxl
  (uses your installed Google Chrome, no browser download needed)

Run
  python amp_price_scraper.py                 # full run
  python amp_price_scraper.py --headful       # watch the browser work
  python amp_price_scraper.py --skip-search   # reuse last shop.app results (faster rerun)
  python amp_price_scraper.py --ship-top 0    # check shipping on EVERY listing (slow)

If a store's domain can't be found automatically, open store_domains.csv,
type the domain next to the store name (example: tfswheels.com), and rerun
with --skip-search.
"""

import argparse
import csv
import json
import os
import random
import re
import statistics
import sys
import time
from collections import defaultdict
from urllib.parse import quote_plus, urlparse

import requests

# ----------------------------------------------------------------- settings
SHIP_ZIP = "60563"            # zip used for shipping quotes (change to a typical customer zip)
SHIP_PROVINCE = "Illinois"
SHIP_COUNTRY = "United States"
UNDERCUT = 1.00               # target price = lowest landed price minus this

# Your own store. Anything matching is shown but left out of "lowest competitor".
OWN_STORE_NAMES = ["wheel tire direct", "naperville wheel", "illini auto tire"]
OWN_DOMAINS = ["mwtakeoffs.com", "naperville-wheel-tire.myshopify.com", "wheeltiredirect"]

MODEL_QUERIES = [
    "AMP Terrain Pro A/T", "AMP Pro AT", "AMP AT Pro",
    "AMP Terrain Attack A/T A", "AMP ATA tire", "AMP Attack AT",
    "AMP Mud Terrain Attack M/T", "AMP MT tire", "AMP Attack MT",
    "AMP Terrain Attack R/T", "AMP RT tire",
    "AMP tires", "AMP tire",
]
RIMS = [15, 16, 17, 18, 20, 22, 24, 26]
RIM_MODELS = ["AMP Pro AT", "AMP ATA", "AMP MT", "AMP RT"]
COMMON_SIZES = ["35x12.50R20", "33x12.50R20", "37x13.50R20", "35x12.50R22", "37x13.50R22",
                "33x12.50R24", "35x12.50R24", "35x12.50R17", "37x12.50R17", "35x12.50R18",
                "285/70R17", "265/70R17", "275/65R20", "275/70R18", "285/65R18", "305/35R24"]

CACHE_DIR = "amp_cache"
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/128.0 Safari/537.36")
IGNORE_HOST_PARTS = ["shop.app", "shopify.com", "shopifycdn", "shopifysvc", "shopifycloud",
                     "google", "gstatic", "facebook", "twitter", "x.com", "instagram",
                     "apple.com", "cloudflare", "jsdelivr", "w3.org", "schema.org",
                     "youtube", "tiktok", "pinterest", "sentry", "doubleclick", "cdn."]

# ----------------------------------------------------------------- parsing helpers
METRIC_RE = re.compile(r"(?<![\d.])(LT|P)?\s*(\d{3})\s*[/\-xX]\s*(\d{2})\s*Z?R\s*-?\s*(\d{2})(?!\d)", re.I)
FLOAT_RE = re.compile(r"(?<![\d.])(\d{2}(?:\.\d)?)\s*[xX]\s*(\d{1,2}(?:\.\d{1,2})?)\s*-?\s*R?\s*(\d{2})(?:\s*LT)?(?![\d.])", re.I)


def parse_size(text):
    """Return a normalized tire size like '35x12.50R20' or '285/70R17', or None."""
    if not text:
        return None
    m = METRIC_RE.search(text)
    if m:
        return f"{m.group(2)}/{m.group(3)}R{m.group(4)}"
    m = FLOAT_RE.search(text)
    if m:
        dia, width, rim = m.group(1), float(m.group(2)), m.group(3)
        if 26 <= float(dia) <= 44 and 8 <= width <= 18:
            dia = dia[:-2] if dia.endswith(".0") else dia
            return f"{dia}x{width:.2f}R{rim}"
    return None


def parse_load(text):
    t = text or ""
    m = re.search(r"\b(?:load\s*range|load|lr)\s*[:\-]?\s*([C-F])\b", t, re.I)
    if m:
        return m.group(1).upper()
    m = re.search(r"\b([C-F])\s*/\s*\d{1,2}\s*PR\b", t, re.I)
    if m:
        return m.group(1).upper()
    m = re.search(r"\b(\d{1,2})\s*PLY\b", t, re.I)
    if m:
        return {"6": "C", "8": "D", "10": "E", "12": "F"}.get(m.group(1), "")
    if re.search(r"\bXL\b", t):
        return "XL"
    return ""


def classify_model(text):
    t = " " + (text or "").lower().replace("-", " ") + " "
    if "r/t" in t or re.search(r"\brt\b", t) or "rugged terrain" in t:
        return "R/T"
    if "m/t" in t or "mud" in t or re.search(r"\bmt\b", t):
        return "M/T"
    if re.search(r"\bpro\b", t):
        return "Terrain Pro A/T P"
    if "attack" in t or re.search(r"\bata\b", t) or "a/t a" in t:
        return "Terrain Attack A/T A"
    if "a/t" in t or re.search(r"\bat\b", t):
        return "A/T (check model)"
    return "Unknown"


def parse_qty(text):
    t = (text or "").lower()
    m = re.search(r"set\s*of\s*(\d+)", t) or re.search(r"\b(\d)\s*(?:pcs|pc|pack|tires)\b", t) \
        or re.search(r"\((\d)\)\s*$", t)
    if m:
        q = int(m.group(1))
        if 1 <= q <= 8:
            return q
    return 1


def looks_like_amp_tire(title, vendor="", product_type="", handle=""):
    blob = f"{title} {vendor} {product_type} {handle}".lower()
    if "amp research" in blob or "power step" in blob or "powerstep" in blob or "bedstep" in blob:
        return False
    if any(w in blob for w in ["package", " kit", "wheel and tire", "wheel & tire", "w/ tires"]):
        return False
    is_amp = (re.search(r"\bamp\b", blob) is not None or "terrain pro a/t" in blob
              or "terrain attack" in blob or "attack m/t" in blob or "amp/ca" in blob)
    return bool(is_amp and parse_size(f"{title} {handle.replace('-', ' ')}"))


def is_own(store_name="", domain=""):
    n, d = (store_name or "").lower(), (domain or "").lower()
    return any(x in n for x in OWN_STORE_NAMES) or any(x in d for x in OWN_DOMAINS)


def parse_card_text(texts):
    """shop.app card innerText -> (store, title, price, compare_at)."""
    best = None
    for t in texts:
        if "$" in t:
            best = t
            break
    if not best:
        return None, (sorted(texts, key=len)[-1] if texts else None), None, None
    lines = [l.strip() for l in best.split("\n") if l.strip()]
    prices = [float(p.replace(",", "")) for l in lines for p in re.findall(r"\$([\d,]+\.\d{2})", l)]
    words = [l for l in lines if "$" not in l and not re.fullmatch(r"\d+% off", l, re.I)
             and not re.fullmatch(r"[\d.]+\s*\(\d+\)", l)]
    store = words[0] if len(words) >= 2 else None
    title = words[1] if len(words) >= 2 else (words[0] if words else None)
    price = prices[0] if prices else None
    compare = prices[1] if len(prices) > 1 and prices[1] > prices[0] else None
    return store, title, price, compare


# ----------------------------------------------------------------- cache helpers
def cpath(name):
    os.makedirs(CACHE_DIR, exist_ok=True)
    return os.path.join(CACHE_DIR, name)


def load_json(name, default):
    try:
        with open(cpath(name), encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return default


def save_json(name, data):
    with open(cpath(name), "w", encoding="utf-8") as f:
        json.dump(data, f, indent=1)


def nap(a=0.6, b=1.4):
    time.sleep(random.uniform(a, b))


def new_session():
    s = requests.Session()
    s.headers.update({"User-Agent": UA, "Accept": "application/json, text/html;q=0.9"})
    return s


def http_get(sess, url, params=None, tries=4):
    for i in range(tries):
        try:
            r = sess.get(url, params=params, timeout=30, allow_redirects=True)
        except requests.RequestException:
            time.sleep(2 * (i + 1))
            continue
        if r.status_code == 429 or r.status_code >= 500:
            time.sleep(5 * (i + 1))
            continue
        return r
    return None


def launch_browser(p, headful):
    """Use your installed Google Chrome first (works on older macOS), then Edge, then Playwright's own."""
    errors = []
    for kw in ({"channel": "chrome"}, {"channel": "msedge"}, {}):
        try:
            return p.chromium.launch(headless=not headful, **kw)
        except Exception as e:
            errors.append(f"{kw or 'bundled chromium'}: {str(e).splitlines()[0]}")
    sys.exit("Could not start a browser. Install Google Chrome (google.com/chrome) and rerun.\n"
             + "\n".join(errors))


# ----------------------------------------------------------------- stage 1: shop.app search
def build_queries():
    q = list(MODEL_QUERIES)
    q += [f"{m} {r}" for m in RIM_MODELS for r in RIMS]
    q += [f"AMP {s}" for s in COMMON_SIZES]
    seen, out = set(), []
    for x in q:
        if x.lower() not in seen:
            seen.add(x.lower())
            out.append(x)
    return out


def scrape_search(page, query, max_rounds=25):
    url = "https://shop.app/search/results?query=" + quote_plus(query)
    page.goto(url, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(3000)
    hits, stale = {}, 0
    for _ in range(max_rounds):
        cards = page.eval_on_selector_all(
            'a[href*="/products/"]',
            "els => els.map(e => ({href: e.href, text: e.innerText || e.getAttribute('aria-label') || ''}))")
        before = len(hits)
        for c in cards:
            m = re.search(r"/products/(\d+)/([^?/#]+)", c["href"])
            if not m:
                continue
            pid, handle = m.group(1), m.group(2)
            v = re.search(r"variantId=(\d+)", c["href"])
            rec = hits.setdefault(pid, {"product_id": pid, "handle": handle,
                                        "variant_id": v.group(1) if v else None,
                                        "url": c["href"].split("#")[0], "texts": []})
            txt = (c["text"] or "").strip()
            if txt and txt not in rec["texts"]:
                rec["texts"].append(txt)
        stale = stale + 1 if len(hits) == before else 0
        if stale >= 3:
            break
        page.mouse.wheel(0, 5000)
        page.wait_for_timeout(1400)
        for label in ("Show more", "Load more", "See more", "More results"):
            try:
                btn = page.get_by_role("button", name=re.compile(label, re.I))
                if btn.count():
                    btn.first.click(timeout=2000)
                    page.wait_for_timeout(1800)
            except Exception:
                pass
    return list(hits.values())


def stage_search(headful):
    from playwright.sync_api import sync_playwright
    queries = build_queries()
    all_hits = load_json("shopapp_hits.json", {})
    done = set(load_json("shopapp_done_queries.json", []))
    with sync_playwright() as p:
        browser = launch_browser(p, headful)
        page = browser.new_page(user_agent=UA, viewport={"width": 1400, "height": 1000})
        for i, q in enumerate(queries, 1):
            if q in done:
                continue
            try:
                hits = scrape_search(page, q)
            except Exception as e:
                print(f"  [{i}/{len(queries)}] {q!r}: failed ({e})")
                continue
            new = 0
            for h in hits:
                store, title, price, compare = parse_card_text(h["texts"])
                h.update(store=store, title=title, price=price, compare_at=compare, query=q)
                h.pop("texts", None)
                if not looks_like_amp_tire(title or "", handle=h["handle"]):
                    continue
                if h["product_id"] not in all_hits:
                    new += 1
                all_hits.setdefault(h["product_id"], h)
            done.add(q)
            save_json("shopapp_hits.json", all_hits)
            save_json("shopapp_done_queries.json", sorted(done))
            print(f"  [{i}/{len(queries)}] {q!r}: {len(hits)} results, {new} new AMP listings "
                  f"(total {len(all_hits)})")
            nap(1.5, 3.0)
        browser.close()
    return all_hits


# ----------------------------------------------------------------- stage 2: store domains
def verify_domain(sess, domain, handle, pid):
    r = http_get(sess, f"https://{domain}/products/{handle}.js", tries=2)
    if not r or r.status_code != 200:
        return None
    try:
        data = r.json()
    except ValueError:
        return None
    if str(data.get("id")) == str(pid):
        return urlparse(r.url).hostname or domain
    return None


def guess_domains(store):
    base = re.sub(r"\b(inc|llc|co|ltd|corp)\b\.?", "", store.lower())
    slug = re.sub(r"[^a-z0-9]", "", base)
    dash = re.sub(r"[^a-z0-9]+", "-", base).strip("-")
    out = [f"{slug}.com", f"www.{slug}.com", f"{dash}.com", f"{slug}.myshopify.com",
           f"{dash}.myshopify.com", f"{slug}.net", f"{slug}.shop", f"{slug}.store", f"{slug}.us"]
    return list(dict.fromkeys(out))


def hosts_in_html(html):
    hosts = set()
    for h in re.findall(r"https?://([a-z0-9][a-z0-9.\-]+\.[a-z]{2,})", html, re.I):
        h = h.lower()
        if not any(x in h for x in IGNORE_HOST_PARTS) or h.endswith(".myshopify.com"):
            hosts.add(h)
    for h in re.findall(r"([a-z0-9\-]+\.myshopify\.com)", html, re.I):
        hosts.add(h.lower())
    return sorted(hosts, key=lambda x: (not x.endswith("myshopify.com"), x))


def stage_domains(hits, headful):
    domains = {}
    if os.path.exists("store_domains.csv"):
        with open("store_domains.csv", newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                if row.get("domain", "").strip():
                    domains[row["store"]] = row["domain"].strip().lower()
    by_store = defaultdict(list)
    for h in hits.values():
        by_store[h.get("store") or "UNKNOWN STORE"].append(h)

    sess = new_session()
    todo = [s for s in by_store if s not in domains]
    need_browser = []
    for store in todo:
        sample = by_store[store][0]
        found = None
        if store != "UNKNOWN STORE":
            for d in guess_domains(store):
                found = verify_domain(sess, d, sample["handle"], sample["product_id"])
                if found:
                    break
        if found:
            domains[store] = found
            print(f"  {store}: {found} (guessed)")
        else:
            need_browser.append(store)

    if need_browser:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = launch_browser(p, headful)
            page = browser.new_page(user_agent=UA)
            for store in need_browser:
                found = None
                for sample in by_store[store][:3]:
                    try:
                        page.goto(sample["url"], wait_until="domcontentloaded", timeout=60000)
                        page.wait_for_timeout(2500)
                        cands = hosts_in_html(page.content())
                        merch = page.eval_on_selector_all(
                            'a[href*="/m/"]', "els => els.map(e => e.href)")
                        if merch:
                            page.goto(merch[0], wait_until="domcontentloaded", timeout=60000)
                            page.wait_for_timeout(2500)
                            cands += hosts_in_html(page.content())
                    except Exception:
                        cands = []
                    for d in list(dict.fromkeys(cands))[:30]:
                        found = verify_domain(sess, d, sample["handle"], sample["product_id"])
                        if found:
                            break
                    if found:
                        break
                    nap()
                if found:
                    domains[store] = found
                    print(f"  {store}: {found} (from shop.app page)")
                else:
                    print(f"  {store}: domain NOT found. Add it to store_domains.csv if you know it.")
            browser.close()

    with open("store_domains.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["store", "domain"])
        for s in sorted(by_store):
            w.writerow([s, domains.get(s, "")])
    return domains, by_store


# ----------------------------------------------------------------- stage 3: store catalogs
def handles_from_html(html):
    return set(re.findall(r'/products/([a-z0-9][a-z0-9\-_%]*)(?=["?#/])', html, re.I))


def collect_handles(sess, domain):
    handles = set()
    sources = []
    for vendor in ("AMP", "AMP Tires", "AMP Tire", "Amp Tires"):
        sources.append((f"https://{domain}/collections/vendors", {"q": vendor}))
    sources.append((f"https://{domain}/search", {"q": "AMP tire", "type": "product"}))
    sources.append((f"https://{domain}/search", {"q": "AMP Terrain", "type": "product"}))
    for url, base in sources:
        for pg in range(1, 60):
            r = http_get(sess, url, params={**base, "page": pg})
            if not r or r.status_code != 200:
                break
            found = handles_from_html(r.text) - handles
            if not found:
                break
            handles |= found
            nap(0.4, 0.9)
    return handles


def product_from_js(data):
    return {
        "id": data.get("id"), "title": data.get("title", ""), "handle": data.get("handle", ""),
        "vendor": data.get("vendor", ""), "product_type": data.get("type", ""),
        "variants": [{"id": v.get("id"), "title": v.get("title", ""),
                      "price": (v.get("price") or 0) / 100.0,
                      "compare_at": (v.get("compare_at_price") or 0) / 100.0 or None,
                      "available": v.get("available")} for v in data.get("variants", [])],
    }


def product_from_json(data):
    return {
        "id": data.get("id"), "title": data.get("title", ""), "handle": data.get("handle", ""),
        "vendor": data.get("vendor", ""), "product_type": data.get("product_type", ""),
        "variants": [{"id": v.get("id"), "title": v.get("title", ""),
                      "price": float(v.get("price") or 0),
                      "compare_at": float(v["compare_at_price"]) if v.get("compare_at_price") else None,
                      "available": v.get("available")} for v in data.get("variants", [])],
    }


def fetch_store_catalog(domain, extra_handles, max_catalog_pages):
    sess = new_session()
    products = {}
    handles = collect_handles(sess, domain) | set(extra_handles)
    for h in sorted(handles):
        r = http_get(sess, f"https://{domain}/products/{h}.js", tries=3)
        if r and r.status_code == 200:
            try:
                p = product_from_js(r.json())
                products[p["id"]] = p
            except ValueError:
                pass
        nap(0.3, 0.7)
    amp = [p for p in products.values()
           if looks_like_amp_tire(p["title"], p["vendor"], p["product_type"], p["handle"])]
    # Fallback: crawl the full catalog if vendor/search pages gave nothing useful.
    if len(amp) <= len(extra_handles) and max_catalog_pages > 0:
        for pg in range(1, max_catalog_pages + 1):
            r = http_get(sess, f"https://{domain}/products.json", params={"limit": 250, "page": pg})
            if not r or r.status_code != 200:
                break
            try:
                batch = r.json().get("products", [])
            except ValueError:
                break
            if not batch:
                break
            for raw in batch:
                p = product_from_json(raw)
                if looks_like_amp_tire(p["title"], p["vendor"], p["product_type"], p["handle"]):
                    products[p["id"]] = p
            nap(0.5, 1.0)
        amp = [p for p in products.values()
               if looks_like_amp_tire(p["title"], p["vendor"], p["product_type"], p["handle"])]
    return amp


def stage_catalogs(domains, by_store, max_catalog_pages):
    listings = []
    cache = load_json("catalogs.json", {})
    for store, recs in sorted(by_store.items()):
        domain = domains.get(store)
        if not domain:
            # No domain: keep the shop.app card price so nothing gets lost.
            for h in recs:
                title = h.get("title") or h["handle"].replace("-", " ")
                qty = parse_qty(title)
                listings.append(make_listing(store, "", h["product_id"], h.get("variant_id"), title, "",
                                             h.get("price"), h.get("compare_at"), None, h["url"],
                                             "shop.app card only (no domain)", qty))
            continue
        if domain not in cache:
            print(f"  pulling AMP catalog from {store} ({domain}) ...")
            cache[domain] = fetch_store_catalog(domain, [h["handle"] for h in recs], max_catalog_pages)
            save_json("catalogs.json", cache)
        products = cache[domain]
        print(f"  {store}: {len(products)} AMP tire products")
        for p in products:
            for v in p["variants"]:
                vt = "" if (v["title"] or "").lower() == "default title" else v["title"]
                qty = parse_qty(f"{p['title']} {vt}")
                listings.append(make_listing(store, domain, p["id"], v["id"], p["title"], vt,
                                             v["price"], v["compare_at"], v["available"],
                                             f"https://{domain}/products/{p['handle']}?variant={v['id']}",
                                             "store site", qty))
    return listings


def make_listing(store, domain, pid, vid, title, vtitle, price, compare, available, url, source, qty):
    full = f"{title} {vtitle}".strip()
    size = parse_size(vtitle) or parse_size(title) or parse_size(url.split("/products/")[-1].replace("-", " "))
    return {
        "store": store, "domain": domain, "own_store": is_own(store, domain),
        "model": classify_model(full), "size": size or "", "load": parse_load(full),
        "title": full, "qty_in_listing": qty,
        "price": price, "compare_at": compare,
        "price_per_tire": round(price / qty, 2) if price else None,
        "available": available, "product_id": pid, "variant_id": vid,
        "url": url, "source": source,
        "ship_qty1": None, "ship_qty4": None, "ship_per_tire": None,
        "landed_per_tire": None, "ship_note": "", "flag": "",
    }


# ----------------------------------------------------------------- stage 4: shipping
def quote_shipping(domain, variant_id, qty):
    sess = new_session()
    r = None
    try:
        r = sess.post(f"https://{domain}/cart/add.js",
                      json={"items": [{"id": int(variant_id), "quantity": qty}]}, timeout=30)
    except requests.RequestException as e:
        return None, f"add to cart error: {e.__class__.__name__}"
    if r.status_code != 200:
        try:
            msg = r.json().get("description") or r.json().get("message") or ""
        except ValueError:
            msg = ""
        return None, f"add to cart failed ({r.status_code}) {msg}".strip()
    params = {"shipping_address[zip]": SHIP_ZIP, "shipping_address[country]": SHIP_COUNTRY,
              "shipping_address[province]": SHIP_PROVINCE}
    rates = None
    r = http_get(sess, f"https://{domain}/cart/shipping_rates.json", params=params, tries=2)
    if r is not None and r.status_code == 200:
        try:
            rates = r.json().get("shipping_rates")
        except ValueError:
            rates = None
    if rates is None:
        try:
            sess.post(f"https://{domain}/cart/prepare_shipping_rates.json", params=params, timeout=30)
            for _ in range(10):
                time.sleep(1.5)
                r = sess.get(f"https://{domain}/cart/async_shipping_rates.json", timeout=30)
                if r.status_code == 200:
                    body = r.json()
                    if body and body.get("shipping_rates") is not None:
                        rates = body["shipping_rates"]
                        break
        except (requests.RequestException, ValueError):
            pass
    if rates is None:
        detail = ""
        try:
            detail = json.dumps(r.json())[:120] if r is not None else ""
        except Exception:
            pass
        return None, f"no rates returned {detail}".strip()
    if not rates:
        return None, "store returned no shipping options for this zip (local pickup only?)"
    best = min(rates, key=lambda x: float(x.get("price") or 0))
    return float(best.get("price") or 0), best.get("name", "")


def group_key(l):
    return (l["model"], l["size"])


def flag_outliers(listings):
    groups = defaultdict(list)
    for l in listings:
        if l["size"] and l["price_per_tire"]:
            groups[group_key(l)].append(l)
    for g in groups.values():
        prices = [l["price_per_tire"] for l in g]
        med = statistics.median(prices)
        for l in g:
            if l["price_per_tire"] > 2.5 * med and len(g) >= 3:
                l["flag"] = "price looks like a set, not per tire"
            elif l["price_per_tire"] < 0.4 * med and len(g) >= 3:
                l["flag"] = "suspiciously low, check listing"


def usable(l):
    return (l["size"] and l["price_per_tire"] and not l["own_store"] and not l["flag"]
            and l["available"] is not False)


def stage_shipping(listings, ship_top):
    cache = load_json("shipping.json", {})
    groups = defaultdict(list)
    for l in listings:
        if usable(l) and l["domain"] and l["variant_id"]:
            groups[group_key(l)].append(l)
    targets = []
    for g in groups.values():
        g.sort(key=lambda x: x["price_per_tire"])
        targets += g if ship_top == 0 else g[:ship_top]
    print(f"  checking shipping on {len(targets)} listings (zip {SHIP_ZIP}) ...")
    for i, l in enumerate(targets, 1):
        per_listing = l["qty_in_listing"] or 1
        units_for_4 = max(1, round(4 / per_listing))   # a "set of 4" listing = 1 unit
        for label, units in (("ship_qty1", 1), ("ship_qty4", units_for_4)):
            if label == "ship_qty4" and units == 1 and per_listing >= 4:
                l[label] = l["ship_qty1"]
                continue
            key = f"{l['domain']}|{l['variant_id']}|{units}|{SHIP_ZIP}"
            if key not in cache:
                cost, note = quote_shipping(l["domain"], l["variant_id"], units)
                cache[key] = [cost, note]
                save_json("shipping.json", cache)
                nap(0.8, 1.6)
            cost, note = cache[key]
            l[label] = cost
            if cost is None:
                l["ship_note"] = (l["ship_note"] + f" | {label}: {note}").strip(" |")
            elif label == "ship_qty1":
                l["ship_note"] = note
        per = None
        if l["ship_qty4"] is not None:
            per = l["ship_qty4"] / (units_for_4 * per_listing)
        elif l["ship_qty1"] is not None:
            per = l["ship_qty1"] / per_listing
        if per is not None:
            l["ship_per_tire"] = round(per, 2)
            l["landed_per_tire"] = round(l["price_per_tire"] + per, 2)
        if i % 10 == 0:
            print(f"    {i}/{len(targets)}")


# ----------------------------------------------------------------- output
def write_excel(listings, domains, by_store, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3A5F")
    money = '"$"#,##0.00'

    def sheet(ws, headers, rows, money_cols=(), widths=None, link_col=None):
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = hdr_font, hdr_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(r)
        for col_idx, h in enumerate(headers, 1):
            letter = get_column_letter(col_idx)
            if h in money_cols:
                for cell in ws[letter][1:]:
                    cell.number_format = money
            ws.column_dimensions[letter].width = (widths or {}).get(h, max(10, min(45, len(h) + 4)))
        if link_col:
            ci = headers.index(link_col) + 1
            for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                for cell in row:
                    if cell.value:
                        cell.hyperlink = cell.value
                        cell.font = Font(color="1155CC", underline="single")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    # Beat Price sheet
    groups = defaultdict(list)
    own = defaultdict(list)
    for l in listings:
        if not l["size"]:
            continue
        if l["own_store"]:
            own[group_key(l)].append(l)
        elif usable(l):
            groups[group_key(l)].append(l)
    rows = []
    keys = sorted(set(groups) | set(own), key=lambda k: (k[0], k[1]))
    for k in keys:
        g = groups.get(k, [])
        stores = sorted({l["store"] for l in g})
        by_sticker = sorted(g, key=lambda x: x["price_per_tire"])
        landed = sorted([l for l in g if l["landed_per_tire"] is not None], key=lambda x: x["landed_per_tire"])
        low_s = by_sticker[0] if by_sticker else None
        low_l = landed[0] if landed else None
        second = landed[1] if len(landed) > 1 else None
        yours = min((l["price_per_tire"] for l in own.get(k, []) if l["price_per_tire"]), default=None)
        target = round(low_l["landed_per_tire"] - UNDERCUT, 2) if low_l else (
            round(low_s["price_per_tire"] - UNDERCUT, 2) if low_s else None)
        if low_l:
            note = ""
            unconf = [l for l in by_sticker if l["landed_per_tire"] is None
                      and l["price_per_tire"] < low_l["landed_per_tire"]]
            if unconf:
                u = unconf[0]
                note = f"{u['store']} is ${u['price_per_tire']:.2f} but shipping unconfirmed, check it"
        elif low_s:
            note = "shipping not confirmed, target based on sticker price"
        else:
            note = "no competitors found"
        rows.append([
            k[0], k[1], len(stores),
            low_s["price_per_tire"] if low_s else None, low_s["store"] if low_s else "",
            low_l["price_per_tire"] if low_l else None, low_l["ship_per_tire"] if low_l else None,
            low_l["landed_per_tire"] if low_l else None, low_l["store"] if low_l else "",
            second["landed_per_tire"] if second else None, second["store"] if second else "",
            yours, target,
            note,
            (low_l or low_s or {}).get("url", ""),
        ])
    ws = wb.active
    ws.title = "Beat Price"
    sheet(ws, ["Model", "Size", "# Stores", "Lowest Sticker", "Lowest Sticker Store",
               "Cheapest Landed: Price", "Ship / Tire", "Landed / Tire", "Cheapest Landed Store",
               "2nd Landed / Tire", "2nd Store", "Your Current Price", "Target Price", "Note", "Link"],
          rows,
          money_cols={"Lowest Sticker", "Cheapest Landed: Price", "Ship / Tire", "Landed / Tire",
                      "2nd Landed / Tire", "Your Current Price", "Target Price"},
          widths={"Model": 20, "Size": 13, "Lowest Sticker Store": 24, "Cheapest Landed Store": 24,
                  "2nd Store": 22, "Note": 40, "Link": 50},
          link_col="Link")
    tgt_col = get_column_letter(13)
    for cell in ws[tgt_col][1:]:
        cell.font = Font(bold=True, color="0B6E2E")

    # All Listings
    cols = ["store", "domain", "own_store", "model", "size", "load", "title", "qty_in_listing",
            "price", "compare_at", "price_per_tire", "ship_qty1", "ship_qty4", "ship_per_tire",
            "landed_per_tire", "available", "flag", "ship_note", "source", "url"]
    ws2 = wb.create_sheet("All Listings")
    srt = sorted(listings, key=lambda l: (l["model"], l["size"], l["price_per_tire"] or 9e9))
    sheet(ws2, cols, [[l.get(c) for c in cols] for l in srt],
          money_cols={"price", "compare_at", "price_per_tire", "ship_qty1", "ship_qty4",
                      "ship_per_tire", "landed_per_tire"},
          widths={"title": 55, "url": 50, "ship_note": 35, "store": 24, "domain": 26, "flag": 30},
          link_col="url")

    # Stores
    ws3 = wb.create_sheet("Stores")
    counts = defaultdict(int)
    for l in listings:
        counts[l["store"]] += 1
    srows = [[s, domains.get(s, ""), counts.get(s, 0), "yes" if is_own(s, domains.get(s, "")) else ""]
             for s in sorted(by_store)]
    sheet(ws3, ["Store", "Domain", "AMP Listings", "Your Store"], srows,
          widths={"Store": 30, "Domain": 32})

    wb.save(out_path)


def listings_from_shopapp(hits):
    out = []
    for h in hits.values():
        title = h.get("title") or h["handle"].replace("-", " ")
        out.append(make_listing(h.get("store") or "UNKNOWN STORE", "", h["product_id"], h.get("variant_id"),
                                title, "", h.get("price"), h.get("compare_at"), None, h["url"],
                                "shop.app", parse_qty(title)))
    return out


def save_outputs(listings, domains, by_store, out_path):
    write_excel(listings, domains, by_store, out_path)
    with open(out_path.replace(".xlsx", ".csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(listings[0].keys()))
        w.writeheader()
        w.writerows(listings)


# ----------------------------------------------------------------- main
def main():
    global SHIP_ZIP
    ap = argparse.ArgumentParser(description="AMP tire competitor pricing from shop.app")
    ap.add_argument("--headful", action="store_true", help="show the browser")
    ap.add_argument("--skip-search", action="store_true", help="reuse cached shop.app results")
    ap.add_argument("--ship-top", type=int, default=5,
                    help="check shipping on the N cheapest listings per model+size (0 = all)")
    ap.add_argument("--max-catalog-pages", type=int, default=40,
                    help="fallback full-catalog crawl limit per store (250 products per page)")
    ap.add_argument("--zip", default=SHIP_ZIP, help="zip code for shipping quotes")
    ap.add_argument("--shopapp-only", action="store_true",
                    help="only use shop.app listing data (fast, no store sites, no shipping check)")
    ap.add_argument("--out", default="amp_prices.xlsx")
    args = ap.parse_args()

    SHIP_ZIP = args.zip

    print("Step 1/4: searching shop.app")
    hits = load_json("shopapp_hits.json", {}) if args.skip_search else stage_search(args.headful)
    if not hits:
        sys.exit("No AMP listings found on shop.app. Try --headful to see what the page shows.")
    print(f"  {len(hits)} AMP listings from shop.app")

    shop_listings = listings_from_shopapp(hits)
    flag_outliers(shop_listings)
    shop_by_store = defaultdict(list)
    for h in hits.values():
        shop_by_store[h.get("store") or "UNKNOWN STORE"].append(h)
    shop_out = args.out if args.shopapp_only else args.out.replace(".xlsx", "_shopapp.xlsx")
    save_outputs(shop_listings, {}, shop_by_store, shop_out)
    print(f"  saved shop.app listings to {shop_out}")
    if args.shopapp_only:
        print(f"\nDone. Open {shop_out} (Beat Price tab first).")
        return

    print("Step 2/4: finding each store's website")
    domains, by_store = stage_domains(hits, args.headful)

    print("Step 3/4: pulling every AMP listing from each store")
    listings = stage_catalogs(domains, by_store, args.max_catalog_pages)
    flag_outliers(listings)
    print(f"  {len(listings)} total listings (variants)")

    print("Step 4/4: checking shipping")
    stage_shipping(listings, args.ship_top)

    save_outputs(listings, domains, by_store, args.out)
    print(f"\nDone. Open {args.out} (Beat Price tab first).")


if __name__ == "__main__":
    main()
